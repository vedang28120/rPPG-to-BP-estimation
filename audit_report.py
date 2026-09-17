"""
audit_report.py
Comprehensive audit of the generated DOCX report.
"""
import docx
from docx.shared import Inches, Pt

doc = docx.Document('LuminaBP_Vocational_Training_Report.docx')

print('='*70)
print('COMPREHENSIVE DOCX AUDIT')
print('='*70)

# 1. Figure captions
print('\n--- FIGURE CAPTIONS ---')
fig_nums = []
for p in doc.paragraphs:
    if p.text.startswith('Figure ') and ':' in p.text[:20]:
        fig_num = p.text.split(':')[0].replace('Figure ', '').strip()
        fig_nums.append(fig_num)
        print(f'  Found: Figure {fig_num} -> {p.text[:90]}')
print(f'  Total figures in captions: {len(fig_nums)}')

# 2. Table captions
print('\n--- TABLE CAPTIONS ---')
tbl_nums = []
for p in doc.paragraphs:
    if p.text.startswith('Table ') and ':' in p.text[:20]:
        tbl_num = p.text.split(':')[0].replace('Table ', '').strip()
        tbl_nums.append(tbl_num)
        print(f'  Found: Table {tbl_num}')
print(f'  Total tables with captions: {len(tbl_nums)}')

# 3. Embedded images
print('\n--- EMBEDDED IMAGES ---')
img_count = 0
for rel in doc.part.rels.values():
    if 'image' in rel.reltype:
        img_count += 1
print(f'  Total embedded images: {img_count}')

# 4. Section settings
print('\n--- SECTION SETTINGS ---')
ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
for i, sect in enumerate(doc.sections):
    sectPr = sect._sectPr
    pgBorders = sectPr.xpath('w:pgBorders')
    pgNumType = sectPr.xpath('w:pgNumType')
    hdr_linked = sect.header.is_linked_to_previous
    ftr_linked = sect.footer.is_linked_to_previous
    
    footer_text = ''
    for fp in sect.footer.paragraphs:
        footer_text += fp.text
    
    print(f'  Section {i}: borders={len(pgBorders)}, pgNumType={len(pgNumType)}, hdr_linked={hdr_linked}, ftr_linked={ftr_linked}, footer="{footer_text[:50]}"')
    if pgNumType:
        fmt = pgNumType[0].get(f'{ns}fmt', 'N/A')
        start = pgNumType[0].get(f'{ns}start', 'N/A')
        print(f'    -> fmt={fmt}, start={start}')

# 5. Content checks
print('\n--- CONTENT CHECKS ---')

# Wrong SOC
for p in doc.paragraphs:
    if 'Snapdragon 8' in p.text and 'Snapdragon 8 Gen 2' in p.text:
        print(f'  *** WRONG SOC: Found "Snapdragon 8 Gen 2" in: "{p.text[:100]}..."')

# Removed persons
for p in doc.paragraphs:
    lower_text = p.text.lower()
    if 'om singh' in lower_text and 'omprakash' not in lower_text:
        print(f'  *** REMOVED PERSON: "Om Singh" found')
    if 'aparna' in lower_text:
        print(f'  *** REMOVED PERSON: "Aparna" found')

# Correct student info
roll_found = any('309302224050' in p.text for p in doc.paragraphs)
enrol_found = any('CE4669' in p.text for p in doc.paragraphs)
batch_found = any('2024 - 2028' in p.text or '2024-2028' in p.text for p in doc.paragraphs)
print(f'  Roll No 309302224050: {"FOUND" if roll_found else "*** MISSING ***"}')
print(f'  Enrollment CE4669: {"FOUND" if enrol_found else "*** MISSING ***"}')
print(f'  Batch 2024-2028: {"FOUND" if batch_found else "*** MISSING ***"}')

# 6. Figure number cross-check against LOF
print('\n--- FIGURE NUMBERING CROSS-CHECK ---')
# Check what the windkessel figure is numbered as in the caption
for p in doc.paragraphs:
    if 'Windkessel' in p.text and 'Figure' in p.text:
        print(f'  Windkessel figure: "{p.text[:100]}"')

# Check Chapter 1 Scope - Snapdragon reference
print('\n--- CHAPTER 1 SCOPE SNAPDRAGON REFERENCE ---')
for p in doc.paragraphs:
    if 'Snapdragon' in p.text:
        print(f'  -> "{p.text[:130]}"')

# 7. Check for signature lines for supervisors (should be removed)
print('\n--- SIGNATURE LINES CHECK ---')
for p in doc.paragraphs:
    if '___' in p.text:
        print(f'  Signature line found: "{p.text[:80]}"')

# 8. Check for "Mr." prefix consistency
print('\n--- TITLE PREFIX CHECK ---')
for p in doc.paragraphs:
    if 'Mr. Vedang' in p.text:
        print(f'  "Mr. Vedang" found in: "{p.text[:80]}"')

# 9. Check for duplicate table borders
print('\n--- TABLE BORDER CHECK ---')
for i, table in enumerate(doc.tables):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = tblPr[0].xpath('w:tblBorders')
        if len(borders) > 1:
            print(f'  *** Table {i}: DUPLICATE borders ({len(borders)} sets)')
    rows = len(table.rows)
    cols = len(table.columns)
    print(f'  Table {i}: {rows}r x {cols}c')

# 10. Check Certificate text for correct phrasing
print('\n--- CERTIFICATE PHRASING ---')
for p in doc.paragraphs:
    if 'studying in 4th semester' in p.text:
        # Check if it says "branch affiliated" (wrong) vs "branch at BIT" (correct)
        if 'branch affiliated' in p.text:
            print(f'  *** WRONG PHRASING: "branch affiliated" found')
            print(f'    -> "{p.text[:200]}"')
        elif 'branch at' in p.text:
            print(f'  Certificate phrasing: CORRECT ("branch at BIT...")')
        else:
            print(f'  Certificate text: "{p.text[:200]}"')

print('\n' + '='*70)
print('AUDIT COMPLETE')
print('='*70)
