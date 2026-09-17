import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_test_doc():
    doc = docx.Document()

    # Section 1: Cover
    s1 = doc.sections[0]
    s1.header.is_linked_to_previous = False
    s1.footer.is_linked_to_previous = False
    doc.add_paragraph('Cover Page')

    # Section 2: Prelim (Roman Lower starting at i)
    s2 = doc.add_section()
    s2.header.is_linked_to_previous = False
    s2.footer.is_linked_to_previous = False
    doc.add_paragraph('Declaration')
    doc.add_page_break()
    doc.add_paragraph('Certificate')

    # Footer for Section 2
    f2 = s2.footer.paragraphs[0]
    f2.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
    r2_txt = f2.add_run('Page | ')
    r2_txt.font.name = 'Times New Roman'
    r2_txt.font.size = docx.shared.Pt(10)
    r2_txt.font.italic = True

    # Use native fldSimple for PAGE
    fld2 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="20"/><w:i/></w:rPr><w:t>i</w:t></w:r></w:fldSimple>')
    f2._p.append(fld2)

    sectPr2 = s2._sectPr
    for ex in sectPr2.xpath('w:pgNumType'):
        sectPr2.remove(ex)
    sectPr2.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1" w:fmt="romanLower"/>'))

    # Section 3: Chapters (Decimal starting at 1)
    s3 = doc.add_section()
    s3.header.is_linked_to_previous = False
    s3.footer.is_linked_to_previous = False
    doc.add_paragraph('Chapter 1')
    doc.add_page_break()
    doc.add_paragraph('Chapter 2')

    f3 = s3.footer.paragraphs[0]
    f3.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
    r3_txt = f3.add_run('Page | ')
    r3_txt.font.name = 'Times New Roman'
    r3_txt.font.size = docx.shared.Pt(10)
    r3_txt.font.italic = True

    fld3 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="20"/><w:i/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>')
    f3._p.append(fld3)

    sectPr3 = s3._sectPr
    for ex in sectPr3.xpath('w:pgNumType'):
        sectPr3.remove(ex)
    sectPr3.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1" w:fmt="decimal"/>'))

    doc.save('test_fldsimple.docx')
    print('test_fldsimple.docx successfully saved!')

if __name__ == '__main__':
    create_test_doc()
