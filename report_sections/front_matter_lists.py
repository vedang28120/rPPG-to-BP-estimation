"""
report_sections/front_matter_lists.py
=====================================
Generates Table of Contents, List of Figures, List of Tables,
and Abbreviations & Nomenclature with clean layouts and EXACT verified page numbers.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def build_table_of_contents(doc, add_heading_chapter, add_table_custom):
    add_heading_chapter(doc, "", "TABLE OF CONTENTS")
    
    headers = ["Section", "Title", "Page No."]
    rows = [
        ["", "Declaration", "i"],
        ["", "Certificate", "ii"],
        ["", "Acknowledgments", "iii"],
        ["", "Abstract", "iv"],
        ["", "List of Figures", "viii"],
        ["", "List of Tables", "ix"],
        ["", "Abbreviations and Nomenclature", "x"],
        ["01", "Introduction", "01"],
        ["1.1", "Background and Clinical Context", "01"],
        ["1.2", "Problem Statement", "02"],
        ["1.3", "Motivation", "03"],
        ["1.4", "Objectives of the Project", "03"],
        ["1.5", "Scope of the Project (Inclusions & Explicit Exclusions)", "04"],
        ["1.6", "Project Contributions", "06"],
        ["1.7", "Organization of the Report", "06"],
        ["02", "Vocational Training Curriculum & Theoretical Foundations", "08"],
        ["2.1", "Scientific Computing Foundations in Python", "08"],
        ["2.2", "Data Preprocessing, Cleaning & Statistical Conditioning", "09"],
        ["2.3", "Machine Learning Taxonomy & Mathematical Core", "11"],
        ["2.4", "Deep Learning Architectures & Optimization", "13"],
        ["2.5", "Advanced AI Frontiers: NLP, Transformers, RAG & Agentic Systems", "15"],
        ["03", "Literature Review & Biomedical Signal Analysis", "18"],
        ["3.1", "Physiological Mechanisms of Blood Pressure Regulation", "18"],
        ["3.2", "Optical Foundations of Remote Photoplethysmography (rPPG)", "19"],
        ["3.3", "Classical and Deep Learning rPPG Extraction Algorithms", "20"],
        ["3.4", "Deep Learning for Optical Blood Pressure Estimation", "21"],
        ["3.5", "Comparative Analysis of Existing Studies", "21"],
        ["3.6", "Research Gaps & The Normotensive Regression Trap", "22"],
        ["3.7", "Proposed LuminaBP Solution", "23"],
        ["04", "Proposed Methodology & System Architecture", "24"],
        ["4.1", "System Overview & Architectural Pipeline", "24"],
        ["4.2", "Module 1: High-Stability Video Acquisition & Exposure Lock", "25"],
        ["4.3", "Module 2: Facial Mesh Landmarking & Dynamic Multi-ROI Tracking", "26"],
        ["4.4", "Module 3: Chrominance Projection (POS Algorithm)", "27"],
        ["4.5", "Module 4: Multi-Stage Signal Conditioning & Savitzky-Golay Denoising", "27"],
        ["4.6", "Module 5: Hemodynamic & Morphological Feature Engineering", "28"],
        ["4.7", "Module 6: Deep Sequential Architecture — MODEL-06-SepHead", "29"],
        ["4.8", "Module 7: Calibration Strategy & Signal Quality Rejection", "30"],
        ["05", "Implementation Details & Experimental Protocol", "32"],
        ["5.1", "Computing Environment, Frameworks, and Libraries", "32"],
        ["5.2", "Hardware Specifications & Mobile Testbeds", "32"],
        ["5.3", "Clinical Benchmark Datasets & Cohort Curation", "33"],
        ["5.4", "Subject-Independent Split Protocol", "34"],
        ["5.5", "Model Training Protocol, Loss Formulations & Optimization", "34"],
        ["5.6", "Mobile Android Application Architecture (LuminaBP Mobile)", "35"],
        ["06", "Experimental Results and Discussion", "36"],
        ["6.1", "Performance Evaluation Standards & Clinical Metrics", "36"],
        ["6.2", "Comparative Model Performance", "36"],
        ["6.3", "Detailed Analysis of Diastolic and Systolic Blood Pressure Estimation", "37"],
        ["6.4", "Clinical Agreement via Bland-Altman Analysis", "38"],
        ["6.5", "Correlation & Error Distribution Analysis", "39"],
        ["6.6", "Comprehensive Ablation Studies", "40"],
        ["6.7", "Hemodynamic Discussion & Physiological Interpretation", "41"],
        ["07", "Conclusion and Future Scope", "43"],
        ["7.1", "Conclusion", "43"],
        ["7.2", "Limitations of the Current Study", "44"],
        ["7.3", "Future Research Directions", "44"],
        ["7.4", "Overall Summary", "45"],
        ["", "References", "46"],
        ["", "Appendix A: Engineering Architecture & Baseline System Specifications", "48"]
    ]
    add_table_custom(doc, "", "", headers, rows, col_widths=[0.65, 4.65, 0.70], show_caption=False, is_prelim=True)

def build_list_of_figures(doc, add_heading_chapter, add_table_custom):
    add_heading_chapter(doc, "", "LIST OF FIGURES")
    
    headers = ["Figure No.", "Figure Name", "Page No."]
    rows = [
        ["3.1", "Physiological Capillary Windkessel Damping & Waveform Morphological Comparison", "19"],
        ["4.1", "End-to-End LuminaBP System Architecture & Processing Pipeline", "24"],
        ["4.2", "MediaPipe 468-Point Facial Mesh & Bilateral Malar/Forehead ROIs", "26"],
        ["4.3", "Multi-Stage Signal Conditioning & Savitzky-Golay Denoising", "28"],
        ["4.4", "MODEL-06-SepHead Deep Neural Network Architecture", "30"],
        ["5.1", "rPPG Signal Extractor Benchmark Performance Comparison", "34"],
        ["6.1", "Bland-Altman Clinical Agreement Analysis on Test Cohort", "38"],
        ["6.2", "Actual vs. Predicted Blood Pressure Scatter Correlation", "39"],
        ["6.3", "Mean Absolute Error (MAE) Boxplot Distribution across Test Cohorts", "40"],
        ["6.4", "MODEL-06-SepHead Diagnostic Regression-to-the-Mean / Template Collapse Analysis", "42"]
    ]
    add_table_custom(doc, "", "", headers, rows, col_widths=[0.85, 4.45, 0.70], show_caption=False, is_prelim=True)

def build_list_of_tables(doc, add_heading_chapter, add_table_custom):
    add_heading_chapter(doc, "", "LIST OF TABLES")
    
    headers = ["Table No.", "Table Name", "Page No."]
    rows = [
        ["2.1", "Machine Learning Performance Metrics Mathematical Summary", "12"],
        ["2.2", "Deep Learning Layer Activations and Mathematical Functions", "13"],
        ["3.1", "Comparative Analysis of State-of-the-Art Optical Blood Pressure Estimation Systems", "21"],
        ["4.1", "LuminaBP Multi-Stage Architectural Pipeline Specifications", "24"],
        ["5.1", "Software Stack and Computational Frameworks", "32"],
        ["5.2", "Hardware Specifications for Training and On-Device Mobile Inference", "33"],
        ["5.3", "MCD Benchmark Dataset Cohort Partitioning & Subject Isolation", "34"],
        ["5.4", "Hyperparameter Configurations for MODEL-06-SepHead", "35"],
        ["6.1", "Performance Comparison of Regression Models on Unseen Test Subjects", "36"],
        ["6.2", "Blood Pressure Estimation Compliance with ISO/AAMI SP10 Standard", "37"],
        ["6.3", "British Hypertension Society (BHS) Standard Grading Evaluation", "38"],
        ["6.4", "Comprehensive Ablation Study on Filtering, ROIs, and Decoupled Heads", "40"],
        ["A.1", "LuminaBP Model Checkpoint Manifest and Parameterization", "48"]
    ]
    add_table_custom(doc, "", "", headers, rows, col_widths=[0.85, 4.45, 0.70], show_caption=False, is_prelim=True)

def build_abbreviations(doc, add_heading_chapter, add_table_custom):
    add_heading_chapter(doc, "", "ABBREVIATIONS AND NOMENCLATURE")
    
    # 4-column compact paired layout to fit on 1 single preliminary page
    headers = ["Abbr.", "Full Meaning / Definition", "Abbr.", "Full Meaning / Definition"]
    rows = [
        ["AAMI", "Assoc. Advancement Medical Instrumentation", "LOSO", "Leave-One-Subject-Out Cross-Validation"],
        ["ACC", "Accelerometer Sensor", "LSTM", "Long Short-Term Memory Network"],
        ["BHS", "British Hypertension Society", "MAE", "Mean Absolute Error"],
        ["BiGRU", "Bidirectional Gated Recurrent Unit", "MAP", "Mean Arterial Pressure (mmHg)"],
        ["BP", "Blood Pressure (Arterial Pressure in mmHg)", "MCD", "Multi-Camera Dataset for rPPG"],
        ["BVP", "Blood Volume Pulse", "MCP", "Model Context Protocol"],
        ["CHROM", "Chrominance-based rPPG Method", "MHSA", "Multi-Head Self-Attention"],
        ["CNN", "Convolutional Neural Network", "MLP", "Multi-Layer Perceptron"],
        ["cPPG", "Contact Photoplethysmography", "NLP", "Natural Language Processing"],
        ["CVD", "Cardiovascular Disease", "POS", "Plane-Orthogonal-to-Skin Algorithm"],
        ["DBP", "Diastolic Blood Pressure (mmHg)", "PWA", "Pulse Wave Analysis"],
        ["DNN", "Deep Neural Network", "PWV", "Pulse Wave Velocity"],
        ["EDA", "Electrodermal Activity", "RAG", "Retrieval-Augmented Generation"],
        ["HR", "Heart Rate (Beats Per Minute, BPM)", "RMSE", "Root Mean Square Error"],
        ["HRV", "Heart Rate Variability", "ROI", "Region of Interest"],
        ["IBI", "Inter-Beat Interval (ms)", "rPPG", "Remote Photoplethysmography"],
        ["ICA", "Independent Component Analysis", "SBP", "Systolic Blood Pressure (mmHg)"],
        ["LLM", "Large Language Model", "SGD", "Stochastic Gradient Descent"],
        ["SNR", "Signal-to-Noise Ratio (dB)", "SVR", "Support Vector Regression"],
        ["TFLite", "TensorFlow Lite Edge Runtime", "TS-CAN", "Temporal Shift Attention Network"]
    ]
    add_table_custom(doc, "", "", headers, rows, col_widths=[0.8, 2.2, 0.8, 2.2], show_caption=False, is_prelim=True)
