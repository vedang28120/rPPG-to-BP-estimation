"""
report_sections/chapter5.py
===========================
Chapter 05: Implementation Details & Experimental Protocol (Updated Hardware Specs)
"""

import docx
from docx.shared import Inches, Pt, RGBColor

def build_chapter5(doc, add_heading_chapter, add_heading_sub1, add_heading_sub2, add_heading_sub3, add_body_p, add_bullet_p, add_figure, add_table_custom):
    add_heading_chapter(doc, "CHAPTER-05", "IMPLEMENTATION DETAILS & EXPERIMENTAL PROTOCOL")
    
    add_body_p(doc, "This chapter provides the technical and operational details of the LuminaBP experimental setup, including the software and hardware frameworks, dataset preparation and curation, subject-isolated partitioning protocols, neural network optimization, and on-device mobile Android engineering.")

    # ---------------------------------------------------------
    # 5.1 Computing Environment, Frameworks, and Libraries
    # ---------------------------------------------------------
    add_heading_sub1(doc, "5.1 Computing Environment, Frameworks, and Libraries")
    add_body_p(doc, "The LuminaBP development workflow was constructed using modular open-source software libraries, ensuring reproducible signal processing, neural network training, and mobile deployment. Table 5.1 details the software environment.")

    t51_headers = ["Software / Library", "Version", "Operational Role in LuminaBP Pipeline"]
    t51_rows = [
        ["Python", "3.10.x", "Primary programming language for research and training pipelines."],
        ["PyTorch / TorchVision", "2.2.0+", "Deep learning framework for MODEL-06-SepHead training and backpropagation."],
        ["TensorFlow / Keras", "2.15.0", "Baseline model evaluation and TFLite model quantization export."],
        ["MediaPipe", "0.10.9", "468-point 3D facial landmark mesh tracking and ROI segmentation."],
        ["OpenCV (cv2)", "4.9.0", "Hardware camera frame capture, colorspace conversions, and image processing."],
        ["SciPy (scipy.signal)", "1.12.0", "Digital Butterworth bandpass filtering and Savitzky-Golay polynomial smoothing."],
        ["Scikit-Learn", "1.4.0", "Statistical regression baselines, feature scaling, and validation metrics."],
        ["Android Studio / SDK", "Iguana / SDK 34", "Native mobile development environment for Android OS."],
        ["TensorFlow Lite (TFLite)", "2.15.0", "Edge inference engine for on-device mobile neural execution."]
    ]
    add_table_custom(doc, "5.1", "Software Stack and Computational Frameworks", t51_headers, t51_rows, col_widths=[1.6, 1.2, 3.6])

    # ---------------------------------------------------------
    # 5.2 Hardware Specifications & Mobile Testbeds
    # ---------------------------------------------------------
    add_heading_sub1(doc, "5.2 Hardware Specifications & Mobile Testbeds")
    add_body_p(doc, "Experimental training and edge deployment were conducted across two dedicated computing environments, detailed in Table 5.2.")

    t52_headers = ["Platform Category", "Hardware Component", "Technical Specification"]
    t52_rows = [
        ["Training Workstation", "Host Processor (CPU)", "Intel Xeon Single Core (2 vCPUs @ 2.20 GHz)"],
        ["Training Workstation", "System Memory (RAM)", "12 GB DDR5 RAM"],
        ["Training Workstation", "Dedicated Accelerator (GPU)", "NVIDIA Tesla T4 (16 GB GDDR6 VRAM, CUDA 12.x)"],
        ["Training Workstation", "High-Speed Storage", "100 GB High-Throughput NVMe Storage"],
        ["Mobile Edge Device", "Mobile SoC / CPU", "Qualcomm Snapdragon 4 Gen 2 (4nm, Octa-Core: 2x 2.2 GHz A78 + 6x 2.0 GHz A55)"],
        ["Mobile Edge Device", "Mobile GPU", "Qualcomm Adreno 613 GPU"],
        ["Mobile Edge Device", "Mobile RAM / OS", "6 GB / 8 GB LPDDR4X Memory / Android 13 / 14 (API Level 33/34)"],
        ["Mobile Edge Device", "Mobile Camera Sensors", "50 MP f/1.8 Primary Sensor / 8 MP f/2.0 Front Sensor (1080p @ 30 FPS, locked AE/AWB)"]
    ]
    add_table_custom(doc, "5.2", "Hardware Specifications for Training and On-Device Mobile Inference", t52_headers, t52_rows, col_widths=[1.5, 1.8, 3.1])

    # ---------------------------------------------------------
    # 5.3 Clinical Benchmark Datasets & Cohort Curation
    # ---------------------------------------------------------
    add_heading_sub1(doc, "5.3 Clinical Benchmark Datasets & Cohort Curation")
    add_body_p(doc, "To establish rigorous clinical validity, LuminaBP was trained and evaluated on the Multi-Camera Dataset for Remote Photoplethysmography (MCD, wengziheng/mcd_rppg). The dataset contains synchronized high-definition facial video streams recorded under varying illumination conditions along with simultaneous clinical ground-truth arterial blood pressure waveforms captured via medical-grade continuous hemodynamic monitors.")
    add_body_p(doc, "A total of 17,943 valid 10-second observation windows across 599 unique human participants were curated. An extensive benchmark of signal extractors was conducted on this cohort, comparing POS against CHROM, ICA, and TS-CAN, as illustrated in Figure 5.1.")

    add_figure(doc, "results/figures/fig_extractor_benchmark_chart.png", "5.1", "rPPG Signal Extractor Benchmark Performance: Comparison of Signal-to-Noise Ratio (SNR) and Heart Rate MAE across POS, CHROM, ICA, and TS-CAN extractors on the MCD dataset.", width_in=5.6)

    # ---------------------------------------------------------
    # 5.4 Subject-Independent Split Protocol
    # ---------------------------------------------------------
    add_heading_sub1(doc, "5.4 Subject-Independent Split Protocol")
    add_body_p(doc, "A major pitfall in biomedical machine learning is subject leakage—where different windows from the same participant appear in both training and test sets, allowing the model to memorize subject-specific facial skin tones rather than learning true physiological dynamics.")
    add_body_p(doc, "To guarantee absolute clinical validity, a strict 100% Subject-Level Partitioning scheme was enforced, as detailed in Table 5.3:")
    add_bullet_p(doc, "Train ∩ Val = ∅, Train ∩ Test = ∅, Val ∩ Test = ∅.", bold_prefix="• ")
    add_bullet_p(doc, "All recording sessions, lighting variations, and windows from any single subject remained strictly confined to one designated fold.", bold_prefix="• ")

    t53_headers = ["Partition Fold", "Subject Count", "Window Count (10s)", "Percentage", "Subject Isolation Status"]
    t53_rows = [
        ["Training Set", "419 Subjects", "12,560 Windows", "70.0%", "Strictly Isolated (0 Overlap)"],
        ["Validation Set", "90 Subjects", "2,692 Windows", "15.0%", "Strictly Isolated (0 Overlap)"],
        ["Testing Set", "90 Subjects", "2,691 Windows", "15.0%", "Strictly Isolated (0 Overlap)"],
        ["Total Cohort", "599 Subjects", "17,943 Windows", "100.0%", "100% Subject-Level Independence"]
    ]
    add_table_custom(doc, "5.3", "MCD Benchmark Dataset Cohort Partitioning & Subject Isolation", t53_headers, t53_rows, col_widths=[1.3, 1.2, 1.4, 0.9, 1.6])

    # ---------------------------------------------------------
    # 5.5 Model Training Protocol, Loss Formulations & Optimization
    # ---------------------------------------------------------
    add_heading_sub1(doc, "5.5 Model Training Protocol, Loss Formulations & Optimization")
    add_body_p(doc, "To mitigate regression-to-the-mean, LuminaBP is trained using a composite multi-objective loss function combining Huber Loss (which is quadratic for small errors and linear for large outliers) with an explicit Pearson Correlation Regularization penalty:")
    add_body_p(doc, "L_{total} = L_{Huber}(y_{sbp}, ŷ_{sbp}) + L_{Huber}(y_{dbp}, ŷ_{dbp}) - λ_r · [r(y_{sbp}, ŷ_{sbp}) + r(y_{dbp}, ŷ_{dbp})]")
    add_body_p(doc, "where λ_r = 0.2 is the correlation weighting factor, enforcing that predicted trajectories track true hemodynamic swings rather than collapsing to static constants.")
    add_body_p(doc, "Table 5.4 outlines the complete hyperparameter configuration employed during model optimization.")

    t54_headers = ["Hyperparameter", "Configured Value", "Technical Rationale"]
    t54_rows = [
        ["Input Dimension", "(1250, 1) @ 125 Hz", "10-second conditioned rPPG pulse sequence."],
        ["Batch Size", "64 Windows", "Balances stochastic gradient noise with GPU parallelization."],
        ["Optimizer", "AdamW", "Decouples weight decay (1e-4) from adaptive gradient moment updates."],
        ["Initial Learning Rate", "1e-3", "Warmup for 5 epochs followed by Cosine Annealing decay to 1e-6."],
        ["BiGRU Hidden Units", "128 (2 Layers)", "Captures bidirectional sequential context."],
        ["MHSA Heads / Dim", "4 Heads / d_k = 32", "Attends to long-range pulse waveform morphology."],
        ["Dropout / LayerNorm", "p = 0.2 / ε = 1e-5", "Regularizes latent feature representations against overfitting."],
        ["Max Epochs / Patience", "150 Epochs / 20 Patience", "Early stopping triggered on validation loss plateau."]
    ]
    add_table_custom(doc, "5.4", "Hyperparameter Configurations for MODEL-06-SepHead", t54_headers, t54_rows, col_widths=[1.6, 1.6, 3.2])

    # ---------------------------------------------------------
    # 5.6 Mobile Android Application Architecture (LuminaBP Mobile)
    # ---------------------------------------------------------
    add_heading_sub1(doc, "5.6 Mobile Android Application Architecture (LuminaBP Mobile)")
    add_body_p(doc, "To realize the vision of ubiquitous cardiovascular assessment, the complete LuminaBP pipeline was engineered into a standalone Android mobile application (LuminaBP Mobile) written in Kotlin.")
    add_body_p(doc, "Key mobile architectural components include:")
    add_bullet_p(doc, "Native Camera2 Pipeline: Background thread capturing 30 FPS camera frames with sensor exposure and white balance lock.", bold_prefix="1. ")
    add_bullet_p(doc, "On-Device MediaPipe Runtime: GPU-accelerated facial mesh detection extracting the 468-landmark wireframe at 30 FPS with <12% CPU utilization.", bold_prefix="2. ")
    add_bullet_p(doc, "TFLite Int8 Quantized Inference: The trained PyTorch model was converted to ONNX and exported to TensorFlow Lite with full 8-bit post-training quantization, shrinking model size from 18.4 MB to 4.2 MB.", bold_prefix="3. ")
    add_bullet_p(doc, "Real-Time HUD Dashboard: Renders dynamic pulse waveforms, real-time heart rate, SBP, DBP, MAP, and ISO/AAMI clinical status overlay.", bold_prefix="4. ")
