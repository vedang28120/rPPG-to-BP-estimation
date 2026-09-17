"""
proofread_report.py
Thorough proofreading of the DOCX.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import docx

doc = docx.Document('LuminaBP_Vocational_Training_Report.docx')

print('='*60)
print('PROOFREADING PASS')
print('='*60)

all_text = []
for p in doc.paragraphs:
    all_text.append(p.text)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            all_text.append(cell.text)

full_text = '\n'.join(all_text)

# 1. Check for common typos and issues
issues = []

# Double spaces
import re
for i, p in enumerate(doc.paragraphs):
    if '  ' in p.text and p.text.strip():
        # Skip signature block (intentional spacing for alignment)
        if '___' not in p.text and 'Signature' not in p.text and "Student" not in p.text and "Roll" not in p.text and "Enrollment" not in p.text:
            issues.append(f'  Double space in para {i}: "...{p.text[max(0,p.text.index("  ")-20):p.text.index("  ")+25]}..."')

# Check for incomplete sentences (ending with : but nothing follows)
# Check for unmatched parentheses
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    open_parens = text.count('(')
    close_parens = text.count(')')
    if open_parens != close_parens and len(text) > 30:
        issues.append(f'  Unmatched parens in para {i} (open={open_parens}, close={close_parens}): "{text[:60]}..."')

# Check for specific factual items
checks = {
    'WHO 17.9 million': '17.9 million' in full_text,
    '1.28 billion': '1.28 billion' in full_text,
    '599 subjects': '599' in full_text,
    '17,943 windows': '17,943' in full_text,
    'DBP MAE 5.91': '5.91' in full_text,
    'SBP MAE 10.12': '10.12' in full_text,
    'ISO/AAMI SP10': 'ISO/AAMI SP10' in full_text or 'ISO/AAMI' in full_text,
    'BHS Grade A': 'Grade A' in full_text,
    'Pearson r = 0.355': '0.355' in full_text,
    'Pearson r = 0.388': '0.388' in full_text,
    'MediaPipe 468': '468' in full_text,
    'POS algorithm': 'Plane-Orthogonal-to-Skin' in full_text,
    'Savitzky-Golay': 'Savitzky-Golay' in full_text,
    'BiGRU': 'BiGRU' in full_text,
    'MHSA': 'MHSA' in full_text,
    'Bland-Altman': 'Bland-Altman' in full_text,
    'TFLite': 'TFLite' in full_text or 'TensorFlow Lite' in full_text,
    'AdamW optimizer': 'AdamW' in full_text,
    'Huber Loss': 'Huber' in full_text,
    'IIIT Naya Raipur': 'IIIT Naya Raipur' in full_text,
    'BIT Raipur': 'Bhilai Institute of Technology' in full_text,
    'CSVTU': 'Chhattisgarh Swami Vivekananda' in full_text,
    'Dr. Anurag Singh': 'Dr. Anurag Singh' in full_text,
    'Prof. OmPrakash Barapatre': 'OmPrakash Barapatre' in full_text,
    'Vedang Bhatt': 'Vedang Bhatt' in full_text,
    'Roll 309302224050': '309302224050' in full_text,
    'Enrollment CE4669': 'CE4669' in full_text,
    'Batch 2024-2028': '2024' in full_text and '2028' in full_text,
    'Intel Xeon': 'Intel Xeon' in full_text,
    'Tesla T4': 'Tesla T4' in full_text,
    '12 GB DDR5': '12 GB DDR5' in full_text,
    'Snapdragon 4 Gen 2': 'Snapdragon 4 Gen 2' in full_text,
    'Adreno 613': 'Adreno 613' in full_text,
    'LPDDR4X': 'LPDDR4X' in full_text,
}

print('\n--- FACTUAL CONTENT CHECK ---')
all_pass = True
for label, found in checks.items():
    status = 'OK' if found else '*** MISSING ***'
    if not found:
        all_pass = False
    print(f'  {label}: {status}')

if all_pass:
    print('\n  [ALL CHECKS PASSED]')

# Print formatting issues
print('\n--- FORMATTING ISSUES ---')
if issues:
    for iss in issues[:20]:
        print(iss)
else:
    print('  No formatting issues detected.')

# Check for "LuminaBP" consistency
lumina_count = full_text.count('LuminaBP')
print(f'\n--- BRAND CONSISTENCY ---')
print(f'  "LuminaBP" mentions: {lumina_count}')

# Check figure/table cross-references in body text
print('\n--- CROSS-REFERENCE CHECK ---')
fig_refs = re.findall(r'Figure [0-9]+\.[0-9]+', full_text)
tbl_refs = re.findall(r'Table [0-9A-Za-z]+\.[0-9]+', full_text)
print(f'  Figure references in text: {len(fig_refs)} -> {sorted(set(fig_refs))}')
print(f'  Table references in text: {len(tbl_refs)} -> {sorted(set(tbl_refs))}')

# Check for any remaining "Snapdragon 8" references
if 'Snapdragon 8' in full_text:
    print('\n  *** WARNING: "Snapdragon 8" still found somewhere!')
else:
    print('\n  Snapdragon SoC references: All correct (4 Gen 2)')

print('\n' + '='*60)
print('PROOFREADING COMPLETE')
print('='*60)
