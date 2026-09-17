"""
build_full_report.py
====================
Complete Report Builder for 'LuminaBP: Non-Invasive Continuous Blood Pressure
Estimation via Remote Photoplethysmography and Deep Learning'.
Configures exact CSVTU/BIT margins, typography, table styling, and footer page numbering.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

HEX_PRIMARY_DARK = "1B4332"  # Deep Forest Teal
HEX_SECONDARY = "0F4C81"     # Classic Medical Blue
HEX_ACCENT = "B8860B"        # Academic Warm Gold
HEX_BG_LIGHT = "F4F6F9"      # Soft Gray Table Header
HEX_BORDER = "CCCCCC"        # Clean Gray Border
HEX_TEXT_MUTED = "555555"    # Subtle Caption Gray

COLOR_BLACK = RGBColor(0, 0, 0)
COLOR_DARK_TEAL = RGBColor(27, 67, 50)
COLOR_MUTED = RGBColor(80, 80, 80)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def prevent_row_splits(table):
    """Ensures table rows do not split across pages and headers repeat."""
    for idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def add_page_border(section):
    sectPr = section._sectPr
    pgBorders = parse_xml(
        f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
        f'<w:top w:val="double" w:sz="12" w:space="24" w:color="1B4332"/>'
        f'<w:bottom w:val="double" w:sz="12" w:space="24" w:color="1B4332"/>'
        f'<w:left w:val="double" w:sz="12" w:space="24" w:color="1B4332"/>'
        f'<w:right w:val="double" w:sz="12" w:space="24" w:color="1B4332"/>'
        f'</w:pgBorders>'
    )
    sectPr.append(pgBorders)

def remove_page_border(section):
    """Explicitly remove any inherited page borders from a section."""
    sectPr = section._sectPr
    for existing in sectPr.xpath('w:pgBorders'):
        sectPr.remove(existing)

def add_footer_page_num(section, is_roman=False, start_val=1):
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""  # Clear existing
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    sectPr = section._sectPr
    for existing in sectPr.xpath('w:pgNumType'):
        sectPr.remove(existing)
        
    fmt_str = "romanLower" if is_roman else "decimal"
    pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="{start_val}" w:fmt="{fmt_str}"/>')
    sectPr.append(pgNumType)
        
    p_run = p.add_run('Page | ')
    p_run.font.name = 'Times New Roman'
    p_run.font.size = Pt(10)
    p_run.font.italic = True
    
    # Use complex field codes (w:fldChar) instead of w:fldSimple for reliable auto-update
    rpr_xml = '<w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="20"/><w:i/></w:rPr>'
    
    # BEGIN field char
    r_begin = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'{rpr_xml}'
        f'<w:fldChar w:fldCharType="begin"/>'
        f'</w:r>'
    )
    p._p.append(r_begin)
    
    # INSTRTEXT with PAGE command
    r_instr = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'{rpr_xml}'
        f'<w:instrText xml:space="preserve"> PAGE </w:instrText>'
        f'</w:r>'
    )
    p._p.append(r_instr)
    
    # SEPARATE field char
    r_sep = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'{rpr_xml}'
        f'<w:fldChar w:fldCharType="separate"/>'
        f'</w:r>'
    )
    p._p.append(r_sep)
    
    # Fallback display text (shown before Word recalculates)
    fallback_txt = "i" if is_roman else "1"
    r_fallback = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'{rpr_xml}'
        f'<w:t>{fallback_txt}</w:t>'
        f'</w:r>'
    )
    p._p.append(r_fallback)
    
    # END field char
    r_end = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'{rpr_xml}'
        f'<w:fldChar w:fldCharType="end"/>'
        f'</w:r>'
    )
    p._p.append(r_end)

def add_heading_chapter(doc, chapter_num_str, chapter_title_str):
    if chapter_num_str:
        p1 = doc.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(16)
        p1.paragraph_format.space_after = Pt(3)
        p1.paragraph_format.keep_with_next = True
        r1 = p1.add_run(chapter_num_str.upper())
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(15)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_BLACK

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(3 if chapter_num_str else 16)
    p2.paragraph_format.space_after = Pt(14)
    p2.paragraph_format.keep_with_next = True
    r2 = p2.add_run(chapter_title_str.upper())
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(15)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_BLACK

def add_heading_sub1(doc, title_str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title_str)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = COLOR_BLACK

def add_heading_sub2(doc, title_str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title_str)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = COLOR_BLACK

def add_heading_sub3(doc, title_str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title_str)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = COLOR_BLACK

def add_body_p(doc, text_str, bold_prefix=None, space_after=5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.45
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
        
    r = p.add_run(text_str)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_BLACK
    return p

def add_bullet_p(doc, text_str, bold_prefix=None, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.left_indent = Inches(0.22 * (level + 1))
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
        
    r = p.add_run(text_str)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_BLACK
    return p

def add_equation_block(doc, eq_lines):
    """
    Renders tightly formatted mathematical equations centered on the page.
    """
    for idx, eq in enumerate(eq_lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2 if idx > 0 else 3)
        p.paragraph_format.space_after = Pt(2 if idx < len(eq_lines)-1 else 5)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = True if idx < len(eq_lines)-1 else False
        
        r = p.add_run(eq)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.5)
        r.font.italic = True
        r.font.color.rgb = COLOR_BLACK

def add_figure(doc, image_path, fig_num_str, fig_caption_str, width_in=5.8):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        p_cap.paragraph_format.keep_with_next = True
        r_cap_lbl = p_cap.add_run(f"Figure {fig_num_str}: ")
        r_cap_lbl.font.name = 'Times New Roman'
        r_cap_lbl.font.size = Pt(10)
        r_cap_lbl.font.bold = True
        r_cap_lbl.font.italic = True
        
        r_cap_txt = p_cap.add_run(fig_caption_str)
        r_cap_txt.font.name = 'Times New Roman'
        r_cap_txt.font.size = Pt(10)
        r_cap_txt.font.italic = True
    else:
        print(f"Warning: Figure image path not found: {image_path}")

def add_table_custom(doc, table_num_str, table_title_str, headers, rows_data, col_widths=None, show_caption=True, is_prelim=False):
    if show_caption and table_num_str:
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(10)
        p_title.paragraph_format.space_after = Pt(3)
        p_title.paragraph_format.keep_with_next = True
        
        r_t_lbl = p_title.add_run(f"Table {table_num_str}: ")
        r_t_lbl.font.name = 'Times New Roman'
        r_t_lbl.font.size = Pt(11.5)
        r_t_lbl.font.bold = True
        
        r_t_txt = p_title.add_run(table_title_str)
        r_t_txt.font.name = 'Times New Roman'
        r_t_txt.font.size = Pt(11.5)
        r_t_txt.font.bold = True
    
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="B0B0B0", sz="4", val="single")
    prevent_row_splits(table)
    
    # Cell vertical margin settings: very compact for prelim tables, clean for standard
    m_top = 20 if is_prelim else 50
    m_bot = 20 if is_prelim else 50
    m_l = 50 if is_prelim else 80
    m_r = 50 if is_prelim else 80
    
    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], HEX_BG_LIGHT)
        set_cell_margins(hdr_cells[i], top=m_top+20, bottom=m_bot+20, left=m_l, right=m_r)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10 if is_prelim else 10.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_BLACK
            
    # Format Body Rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        is_bold_row = is_prelim and len(row_values) > 1 and row_values[0].isdigit() and len(row_values[0]) == 2
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = str(val)
            set_cell_margins(row_cells[c_idx], top=m_top, bottom=m_bot, left=m_l, right=m_r)
            if not is_prelim and r_idx % 2 == 1:
                set_cell_background(row_cells[c_idx], "FAFAFA")
            p = row_cells[c_idx].paragraphs[0]
            if is_prelim:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 or len(str(val)) > 20 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(0.5)
            p.paragraph_format.line_spacing = 1.05
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9.5 if is_prelim else 10)
                if is_bold_row:
                    r.font.bold = True
                r.font.color.rgb = COLOR_BLACK
                
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
                
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4 if is_prelim else 6)
